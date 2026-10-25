"""Attendance acceptance scenarios (scope doc, Attendance). Real tenant schema.

    python manage.py test apps.attendance.tests.test_attendance
"""
import io
from datetime import date

from django.core.files.uploadedfile import SimpleUploadedFile
from django_tenants.test.client import TenantClient
from django_tenants.utils import schema_context
from freezegun import freeze_time
from openpyxl import load_workbook

from apps.attendance import services
from apps.attendance.models import AttendanceCard, AttendanceDevice, AttendanceRecord, AttendanceSettings, Holiday
from apps.core.exceptions import Conflict
from apps.core.models import AuditLog
from apps.core.testing import EraTestCase
from apps.people.models import Enrollment, Program, Section, Subject, TeachingAssignment


class AttendanceBase(EraTestCase):
    def setUp(self):
        super().setUp()
        self.program = Program.objects.create(branch=self.branch, code="C8", name="Class 8")
        self.teacher_staff = self.make_staff()
        self.teacher = self.make_user("teacher")
        self.teacher_staff.user = self.teacher
        self.teacher_staff.save()
        self.section = Section.objects.create(branch=self.branch, academic_year=self.year, program=self.program, name="A",
                                              class_teacher=self.teacher_staff)
        self.other = Section.objects.create(branch=self.branch, academic_year=self.year, program=self.program, name="B")
        self.students = [self.enrol(self.section, i) for i in range(1, 41)]

    def enrol(self, section, roll):
        s = self.make_student()
        Enrollment.objects.create(branch=self.branch, student=s, academic_year=self.year, program=self.program,
                                  section=section, roll_no=str(roll))
        return s

    def mark(self, user, on, absent=(), statuses=None, section=None):
        return self.api(user).post("/api/attendance/records/mark/", {
            "section": (section or self.section).pk, "date": str(on), "absent": [s.pk for s in absent],
            "statuses": {str(k.pk): v for k, v in (statuses or {}).items()},
        }, content_type="application/json")


class MarkingTest(AttendanceBase):
    @freeze_time("2026-10-26 09:00:00+05:30")  # a Monday
    def test_class_of_forty_with_three_absentees_in_one_save(self):
        r = self.mark(self.teacher, date(2026, 10, 26), absent=self.students[:3])
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()["created"], 40)
        qs = AttendanceRecord.objects.filter(date=date(2026, 10, 26), section=self.section)
        self.assertEqual((qs.count(), qs.filter(status="absent").count()), (40, 3))

    def test_same_day_correction_by_teacher_past_correction_only_by_admin(self):
        with freeze_time("2026-10-19 09:00:00+05:30"):
            self.assertEqual(self.mark(self.teacher, date(2026, 10, 19), absent=self.students[:2]).status_code, 200)
            r = self.mark(self.teacher, date(2026, 10, 19), absent=self.students[:1])  # same-day fix
            self.assertEqual((r.status_code, r.json()["updated"]), (200, 1))
        with freeze_time("2026-10-26 09:00:00+05:30"):
            r = self.mark(self.teacher, date(2026, 10, 19), absent=[])
            self.assertEqual(r.status_code, 409)
            self.assertIn("Admin", r.json()["detail"])
            r = self.mark(self.make_user("admin"), date(2026, 10, 19), absent=[])
            self.assertEqual((r.status_code, r.json()["updated"]), (200, 1))
        with schema_context("public"):
            log = AuditLog.objects.filter(action="attendance.corrected").order_by("-id").first()
        self.assertEqual(log.detail["changes"], [{"student": self.students[0].pk, "from": "absent", "to": "present"}])

    @freeze_time("2026-10-26 09:00:00+05:30")
    def test_teacher_cannot_mark_another_section_and_statuses_apply(self):
        b = self.enrol(self.other, 1)
        self.assertEqual(self.mark(self.teacher, date(2026, 10, 26), section=self.other).status_code, 403)
        r = self.mark(self.make_user("admin"), date(2026, 10, 26), section=self.other, statuses={b: "late"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(AttendanceRecord.objects.get(student=b).status, "late")
        # a subject teacher of section B may mark it
        subject = Subject.objects.create(branch=self.branch, code="M", name="Maths", program=self.program)
        TeachingAssignment.objects.create(branch=self.branch, staff=self.teacher_staff, subject=subject, section=self.other)
        self.assertEqual(self.mark(self.teacher, date(2026, 10, 26), section=self.other).status_code, 200)

    def test_future_date_holiday_and_per_lecture_rules(self):
        with freeze_time("2026-10-26 09:00:00+05:30"):
            with self.assertRaises(Conflict):
                services.mark_section(section=self.section, on=date(2026, 10, 27), statuses={}, by="x", can_correct_past=True)
            with self.assertRaisesMessage(Conflict, "holiday"):
                services.mark_section(section=self.section, on=date(2026, 10, 25), statuses={}, by="x", can_correct_past=True)  # Sunday
            AttendanceSettings.objects.update_or_create(branch=self.branch, defaults={"mode": "per_lecture"})
            with self.assertRaisesMessage(Conflict, "lecture"):
                services.mark_section(section=self.section, on=date(2026, 10, 26), statuses={}, by="x", can_correct_past=True)
            services.mark_section(section=self.section, on=date(2026, 10, 26), period=1, statuses={}, by="x", can_correct_past=True)
            services.mark_section(section=self.section, on=date(2026, 10, 26), period=2, statuses={}, by="x", can_correct_past=True)
        self.assertEqual(AttendanceRecord.objects.filter(student=self.students[0]).count(), 2)


class ReportTest(AttendanceBase):
    def test_monthly_register_leaves_holidays_blank_and_out_of_percentage(self):
        Holiday.objects.create(branch=self.branch, date=date(2026, 10, 2), name="Gandhi Jayanti")
        s = self.students[0]
        with freeze_time("2026-10-31 18:00:00+05:30"):
            for day in (1, 3, 5, 6):  # 4 working days marked; 2 Oct holiday, 4 Oct Sunday
                services.mark_section(section=self.section, on=date(2026, 10, day),
                                      absent=[s.pk] if day == 6 else [], statuses={}, by="t", can_correct_past=True)
        data = services.register(section=self.section, month=date(2026, 10, 1))
        row = next(r for r in data["rows"] if r["student"] == s.pk)
        self.assertIsNone(row["marks"][1])  # 2 Oct
        self.assertIsNone(row["marks"][3])  # 4 Oct (Sunday)
        self.assertEqual(row["marks"][0], "P")
        self.assertEqual((row["attended"], row["marked"], row["percentage"]), (3, 4, 75.0))
        self.assertEqual(data["working_days"], 31 - 4 - 1)  # 4 Sundays + Gandhi Jayanti

        from apps.attendance.pdf import register_document

        self.assertLessEqual(len(register_document(data, branch=self.branch).pages), 2)
        r = self.api(self.make_user("admin")).get(f"/api/attendance/reports/register/?section={self.section.pk}&month=2026-10&pdf=1")
        self.assertEqual(r["Content-Type"], "application/pdf")

    def test_students_below_75_for_the_term_export_to_excel(self):
        weak, ok = self.students[0], self.students[1]
        with freeze_time("2026-10-31 18:00:00+05:30"):
            for day in (5, 6, 7, 8):
                services.mark_section(section=self.section, on=date(2026, 10, day),
                                      absent=[weak.pk] if day in (6, 7) else [], statuses={}, by="t", can_correct_past=True)
        api = self.api(self.make_user("admin"))
        rows = api.get("/api/attendance/reports/below/?from=2026-10-01&to=2026-10-31").json()
        self.assertEqual([r["student"] for r in rows], [weak.pk])
        self.assertEqual(rows[0]["percentage"], 50.0)
        self.assertNotIn(ok.pk, [r["student"] for r in rows])
        x = api.get("/api/attendance/reports/below/?from=2026-10-01&to=2026-10-31&format=xlsx")
        ws = load_workbook(io.BytesIO(x.content)).active
        self.assertEqual(ws.max_row, 2)
        self.assertIn("percentage", [c.value for c in ws[1]])

    def test_long_absentees_and_daily_summary(self):
        s = self.students[0]
        with freeze_time("2026-10-31 18:00:00+05:30"):
            for day in (22, 23, 24, 26, 27, 28, 29, 30):  # skips Sunday 25
                services.mark_section(section=self.section, on=date(2026, 10, day),
                                      absent=[s.pk] if day >= 24 else [], statuses={}, by="t", can_correct_past=True)
        rows = services.long_absentees(branch=self.branch, as_of=date(2026, 10, 31), min_days=5)
        self.assertEqual([(r["student"], r["absent_days"]) for r in rows], [(s.pk, 6)])
        summary = {r["section"]: r for r in services.daily_summary(branch=self.branch, on=date(2026, 10, 30))}
        self.assertEqual((summary[self.section.pk]["present"], summary[self.section.pk]["absent"]), (39, 1))


class DeviceTest(AttendanceBase):
    def setUp(self):
        super().setUp()
        self.device = AttendanceDevice.objects.create(branch=self.branch, name="Gate", brand="zkteco", serial_no="CQZ7231")
        self.card_student = self.students[5]
        AttendanceCard.objects.create(branch=self.branch, code="1006", student=self.card_student)

    @freeze_time("2026-10-26 08:05:00+05:30")
    def test_rfid_tap_marks_present_at_once_and_no_punch_stays_unmarked(self):
        c = TenantClient(self.tenant)
        r = c.get("/iclock/cdata?SN=CQZ7231&options=all")
        self.assertEqual(r.status_code, 200)
        self.assertIn("GET OPTION FROM: CQZ7231", r.content.decode())
        r = c.post("/iclock/cdata?SN=CQZ7231&table=ATTLOG&Stamp=1", data="1006\t2026-10-26 08:01:12\t0\t1\t0\t0\n9999\t2026-10-26 08:02:00\t0\t1\n",
                   content_type="text/plain")
        self.assertEqual(r.content.decode(), "OK: 2")
        rec = AttendanceRecord.objects.get(student=self.card_student, date=date(2026, 10, 26))
        self.assertEqual((rec.status, rec.source, rec.section), ("present", "device", self.section))
        self.assertFalse(AttendanceRecord.objects.filter(student=self.students[6]).exists())  # no punch → unmarked

        sheet = self.api(self.teacher).get(f"/api/attendance/records/sheet/?section={self.section.pk}&date=2026-10-26").json()
        row = next(x for x in sheet["students"] if x["student"] == self.card_student.pk)
        self.assertEqual((row["status"], row["punched"]), ("present", True))
        # the teacher's save keeps the device mark and fills in the rest
        self.mark(self.teacher, date(2026, 10, 26), absent=self.students[6:7])
        self.assertEqual(AttendanceRecord.objects.filter(date=date(2026, 10, 26)).count(), 40)
        # resend of the same log is ignored
        c.post("/iclock/cdata?SN=CQZ7231&table=ATTLOG", data="1006\t2026-10-26 08:01:12\t0\t1\n", content_type="text/plain")
        self.assertEqual(self.card_student.punches.count(), 1)

    def test_unknown_or_inactive_device_refused(self):
        c = TenantClient(self.tenant)
        self.assertEqual(c.get("/iclock/cdata?SN=NOPE").status_code, 403)
        self.device.is_active = False
        self.device.save()
        self.assertEqual(c.get("/iclock/cdata?SN=CQZ7231").status_code, 403)

    @freeze_time("2026-10-26 10:00:00+05:30")
    def test_json_push_needs_the_device_key(self):
        generic = AttendanceDevice.objects.create(branch=self.branch, name="Bridge", brand="generic", serial_no="BR-1")
        c = TenantClient(self.tenant)
        body = {"punches": [{"code": "1006", "time": "2026-10-26 08:00:00"}]}
        self.assertEqual(c.post("/api/attendance/push/", body, content_type="application/json").status_code, 403)
        r = c.post("/api/attendance/push/", body, content_type="application/json", HTTP_X_DEVICE_KEY=generic.token)
        self.assertEqual(r.json()["stored"], 1)

    @freeze_time("2026-10-27 10:00:00+05:30")
    def test_import_any_devices_punch_log(self):
        staff = self.make_staff()
        AttendanceCard.objects.create(branch=self.branch, code="E-77", staff=staff)
        csv_ok = "User ID,Date,Time\n1006,27-10-2026,07:58:31\nE-77,27-10-2026,08:10:00\n"
        api = self.api(self.make_user("admin"))
        r = api.post("/api/attendance/punches/import/", {"file": SimpleUploadedFile("log.csv", csv_ok.encode()), "dry_run": "true"})
        self.assertEqual((r.status_code, r.json()["saved"]), (200, 0))
        r = api.post("/api/attendance/punches/import/", {"file": SimpleUploadedFile("log.csv", csv_ok.encode())})
        self.assertEqual(r.json()["saved"], 2, r.content)
        self.assertEqual(staff.punches.count(), 1)  # kept for HR attendance
        self.assertTrue(AttendanceRecord.objects.filter(student=self.card_student, date=date(2026, 10, 27), source="device").exists())
        bad = "User ID,Time\n1006,yesterday\n"
        r = api.post("/api/attendance/punches/import/", {"file": SimpleUploadedFile("bad.csv", bad.encode())})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["errors"][0]["row"], 2)
        r = self.api(self.teacher).post("/api/attendance/punches/import/", {"file": SimpleUploadedFile("log.csv", csv_ok.encode())})
        self.assertEqual(r.status_code, 403)
