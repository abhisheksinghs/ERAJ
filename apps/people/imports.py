"""
Bulk import of students and staff from Excel/CSV (scope: "row-level errors
shown before anything is saved").

The whole file is validated first. If any row fails, nothing is written and
every error is returned with its row number; with `dry_run` nothing is written
even when all rows pass. Writes happen in one transaction.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime

from django.db import transaction
from openpyxl import load_workbook

from apps.people.models import (
    AcademicYear,
    Department,
    Enrollment,
    Guardian,
    Program,
    Section,
    Staff,
    Student,
    StudentGuardian,
)

STUDENT_COLUMNS = [
    "admission_no", "full_name", "name_hi", "date_of_birth", "gender", "phone", "email", "address",
    "program_code", "section", "roll_no", "guardian_name", "guardian_phone", "guardian_relation",
]
STAFF_COLUMNS = [
    "employee_code", "full_name", "name_hi", "designation", "department", "phone", "email",
    "date_of_birth", "date_joined", "employment_type",
]
MAX_ROWS = 5000


class ImportFileError(Exception):
    pass


def read_rows(upload) -> list[dict]:
    name = (upload.name or "").lower()
    if name.endswith(".xlsx"):
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    elif name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise ImportFileError("Upload an .xlsx or .csv file.")
    if not rows:
        raise ImportFileError("The file is empty.")
    header = [str(h or "").strip().lower() for h in rows[0]]
    body = [r for r in rows[1:] if any(c not in (None, "") for c in r)]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"At most {MAX_ROWS} rows per file.")
    return [
        {header[i]: ("" if v is None else v) for i, v in enumerate(r) if i < len(header)}
        for r in body
    ]


def _text(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(v):
    if v in ("", None):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(_text(v), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{v}' is not a date (use YYYY-MM-DD or DD-MM-YYYY)")


def import_students(*, branch, rows: list[dict], dry_run: bool) -> dict:
    year = AcademicYear.objects.filter(is_current=True).first()
    programs = {p.code.lower(): p for p in Program.objects.filter(branch=branch)}
    existing = set(Student.all_objects.filter(branch=branch).values_list("admission_no", flat=True))
    errors, valid, seen = [], [], set()
    genders = {c for c, _ in Student.Gender.choices}
    relations = {c for c, _ in StudentGuardian.Relation.choices}

    for n, row in enumerate(rows, start=2):  # row 1 is the header
        r = {k: _text(row.get(k, "")) for k in STUDENT_COLUMNS}
        problems = []
        if not r["admission_no"]:
            problems.append("admission_no is required")
        elif r["admission_no"] in existing or r["admission_no"] in seen:
            problems.append(f"admission_no {r['admission_no']} already exists")
        if not r["full_name"]:
            problems.append("full_name is required")
        if r["gender"] and r["gender"].lower() not in genders:
            problems.append(f"gender must be one of {sorted(genders)}")
        try:
            dob = _date(row.get("date_of_birth"))
        except ValueError as exc:
            problems.append(str(exc))
            dob = None
        program = programs.get(r["program_code"].lower()) if r["program_code"] else None
        if r["program_code"] and program is None:
            problems.append(f"program_code {r['program_code']} not found")
        if program and year is None:
            problems.append("set a current academic year before enrolling students")
        if r["guardian_name"] and not r["guardian_phone"]:
            problems.append("guardian_phone is required when guardian_name is given")
        if r["guardian_relation"] and r["guardian_relation"].lower() not in relations:
            problems.append(f"guardian_relation must be one of {sorted(relations)}")
        if problems:
            errors.append({"row": n, "errors": problems})
            continue
        seen.add(r["admission_no"])
        valid.append((r, dob, program))

    if errors or dry_run:
        return {"created": 0 if errors else len(valid), "errors": errors, "dry_run": dry_run or bool(errors)}

    with transaction.atomic():
        for r, dob, program in valid:
            student = Student.objects.create(
                branch=branch, admission_no=r["admission_no"], full_name=r["full_name"], name_hi=r["name_hi"],
                date_of_birth=dob, gender=r["gender"].lower(), phone=r["phone"], email=r["email"],
                address=r["address"],
            )
            if program:
                section = None
                if r["section"]:
                    section, _ = Section.objects.get_or_create(
                        branch=branch, academic_year=year, program=program, term=None, name=r["section"]
                    )
                Enrollment.objects.create(
                    branch=branch, student=student, academic_year=year, program=program,
                    section=section, roll_no=r["roll_no"],
                )
            if r["guardian_name"]:
                guardian, _ = Guardian.objects.get_or_create(
                    branch=branch, phone=r["guardian_phone"], defaults={"full_name": r["guardian_name"]}
                )
                StudentGuardian.objects.create(
                    branch=branch, student=student, guardian=guardian,
                    relation=(r["guardian_relation"] or "guardian").lower(), is_primary=True,
                )
    return {"created": len(valid), "errors": [], "dry_run": False}


def import_staff(*, branch, rows: list[dict], dry_run: bool) -> dict:
    existing = set(Staff.all_objects.filter(branch=branch).values_list("employee_code", flat=True))
    types = {c for c, _ in Staff.EmploymentType.choices}
    errors, valid, seen = [], [], set()
    for n, row in enumerate(rows, start=2):
        r = {k: _text(row.get(k, "")) for k in STAFF_COLUMNS}
        problems = []
        if not r["employee_code"]:
            problems.append("employee_code is required")
        elif r["employee_code"] in existing or r["employee_code"] in seen:
            problems.append(f"employee_code {r['employee_code']} already exists")
        if not r["full_name"]:
            problems.append("full_name is required")
        if r["employment_type"] and r["employment_type"].lower() not in types:
            problems.append(f"employment_type must be one of {sorted(types)}")
        dates = {}
        for col in ("date_of_birth", "date_joined"):
            try:
                dates[col] = _date(row.get(col))
            except ValueError as exc:
                problems.append(f"{col}: {exc}")
        if problems:
            errors.append({"row": n, "errors": problems})
            continue
        seen.add(r["employee_code"])
        valid.append((r, dates))

    if errors or dry_run:
        return {"created": 0 if errors else len(valid), "errors": errors, "dry_run": dry_run or bool(errors)}

    with transaction.atomic():
        for r, dates in valid:
            dept = None
            if r["department"]:
                dept, _ = Department.objects.get_or_create(branch=branch, name=r["department"])
            Staff.objects.create(
                branch=branch, employee_code=r["employee_code"], full_name=r["full_name"], name_hi=r["name_hi"],
                designation=r["designation"], department=dept, phone=r["phone"], email=r["email"],
                date_of_birth=dates["date_of_birth"],
                **({"date_joined": dates["date_joined"]} if dates["date_joined"] else {}),
                employment_type=(r["employment_type"] or "permanent").lower(),
            )
    return {"created": len(valid), "errors": [], "dry_run": False}
