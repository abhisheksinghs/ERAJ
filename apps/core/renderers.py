"""
Excel and CSV for every list endpoint (scope: "every list screen exports to
Excel and CSV"):  GET /api/<module>/<resource>/?format=xlsx  (or csv)

Exports skip pagination (see StandardPagination) and are capped at
EXPORT_MAX_ROWS. Nested objects are flattened: {"primary_guardian": {"phone": …}}
becomes a "primary_guardian.phone" column.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime

from rest_framework.renderers import BaseRenderer

EXPORT_MAX_ROWS = 20000

# Spreadsheet formula injection (OWASP "CSV injection"): text a user typed —
# a name, an address — that starts with one of these is run as a formula when
# the export is opened in Excel. Such text is exported with a leading
# apostrophe; real numbers (e.g. "-500.00" credit notes) are left as numbers.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _is_number(text: str) -> bool:
    try:
        float(text)
        return True
    except ValueError:
        return False


def safe_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_START) and not _is_number(value):
        return "'" + value
    return value


def _flatten(row: dict, prefix: str = "") -> dict:
    out = {}
    for k, v in row.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(_flatten(v, f"{key}."))
        elif isinstance(v, list):
            out[key] = ", ".join(str(x) for x in v)
        else:
            out[key] = v
    return out


def _rows(data) -> list[dict]:
    if isinstance(data, dict) and "results" in data:
        data = data["results"]
    if isinstance(data, dict):
        data = [data]
    return [_flatten(r) for r in (data or [])[:EXPORT_MAX_ROWS] if isinstance(r, dict)]


def _columns(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def _filename(renderer_context, ext: str) -> str:
    view = renderer_context.get("view")
    base = getattr(view, "basename", None) or "export"
    return f"{base}-{datetime.now():%Y%m%d-%H%M}.{ext}"


class XlsxRenderer(BaseRenderer):
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    format = "xlsx"
    charset = None
    render_style = "binary"

    def render(self, data, accepted_media_type=None, renderer_context=None):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        renderer_context = renderer_context or {}
        response = renderer_context.get("response")
        wb = Workbook()
        ws = wb.active
        ws.title = "Export"
        if response is not None and response.status_code >= 400:
            ws.append(["error"])
            ws.append([str(data)])
        else:
            rows = _rows(data)
            cols = _columns(rows)
            ws.append(cols)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for r in rows:
                ws.append([safe_cell(r.get(c) if not isinstance(r.get(c), (dict, list)) else str(r.get(c))) for c in cols])
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == "f":  # openpyxl turns any "=…" string into a formula
                        cell.data_type = "s"
            ws.freeze_panes = "A2"
            for i, c in enumerate(cols, start=1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(40, max(10, len(c) + 2))
        if response is not None:
            response["Content-Disposition"] = f'attachment; filename="{_filename(renderer_context, "xlsx")}"'
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


class CsvRenderer(BaseRenderer):
    media_type = "text/csv"
    format = "csv"
    charset = "utf-8"

    def render(self, data, accepted_media_type=None, renderer_context=None):
        renderer_context = renderer_context or {}
        response = renderer_context.get("response")
        rows = _rows(data) if not (response is not None and response.status_code >= 400) else [{"error": str(data)}]
        cols = _columns(rows)
        buf = io.StringIO()
        buf.write("﻿")  # BOM so Excel opens Hindi text as UTF-8
        writer = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({k: safe_cell(v) for k, v in r.items()} for r in rows)
        if response is not None:
            response["Content-Disposition"] = f'attachment; filename="{_filename(renderer_context, "csv")}"'
        return buf.getvalue().encode("utf-8")
