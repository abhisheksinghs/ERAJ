"""Exports must not turn user-typed text into spreadsheet formulas."""
import io

from django.test import SimpleTestCase
from openpyxl import load_workbook

from apps.core.renderers import CsvRenderer, XlsxRenderer, safe_cell


class FormulaInjectionTest(SimpleTestCase):
    ROWS = {"results": [{"name": '=HYPERLINK("http://evil","x")', "amount": "-500.00", "note": "@SUM(A1)"}]}

    def test_safe_cell(self):
        self.assertEqual(safe_cell("=1+1"), "'=1+1")
        self.assertEqual(safe_cell("+91 98000"), "'+91 98000")
        self.assertEqual(safe_cell("-500.00"), "-500.00")  # a number stays a number
        self.assertEqual(safe_cell("Ravi"), "Ravi")
        self.assertEqual(safe_cell(5), 5)

    def test_xlsx_has_no_formulas(self):
        wb = load_workbook(io.BytesIO(XlsxRenderer().render(self.ROWS, renderer_context={})))
        cells = [c for row in wb.active.iter_rows(min_row=2) for c in row]
        self.assertFalse([c.coordinate for c in cells if c.data_type == "f"])
        self.assertTrue(str(cells[0].value).startswith("'="))

    def test_csv_escapes_formula_text(self):
        text = CsvRenderer().render(self.ROWS, renderer_context={}).decode("utf-8")
        self.assertIn("'=HYPERLINK", text)
        self.assertIn("'@SUM", text)
        self.assertIn(",-500.00,", text)
